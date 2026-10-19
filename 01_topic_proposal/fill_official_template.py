"""Fill the official NCKH FA26 RESEARCH PROPOSAL TEMPLATE.docx with the proposal content, keeping the template's own
formatting, headings, numbering and commitment wording.

The content (text, references, plan, appendix, AI disclosure) is read from build_nckh_proposal.py so both outputs stay
identical in substance. Guidance paragraphs of the template are replaced by content; the closing "Note" (formatting
instructions for authors) is removed.

Usage:
    python 01_topic_proposal/fill_official_template.py "<template.docx>" 01_topic_proposal/NCKH_FA26_Proposal_official.docx
"""
import copy
import io
import os
import re
import sys

import docx
from docx.shared import Pt
from docx.text.paragraph import Paragraph

TPL, OUT = sys.argv[1], sys.argv[2]

# ---------------------------------------------------------------- load content from the builder (without building)
src = io.open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "build_nckh_proposal.py"), encoding="utf-8").read()
src = src.split("# ----------------------------------------------------------------------------------------- numbering")[0]
ns = {}
_argv = sys.argv
sys.argv = ["x", "unused.docx"]
exec(compile(src, "build_nckh_proposal.py", "exec"), ns)
sys.argv = _argv
C, REFS, TITLE = ns["C"], ns["REFS"], ns["TITLE"]

order = []


def cite(text):
    def rep(m):
        k = m.group(1)
        if k not in order:
            order.append(k)
        return f"[{order.index(k) + 1}]"
    return re.sub(r"\{(\w+)\}", rep, text)


# group the content by section heading
sections, cur = {}, None
for kind, val in C:
    if kind in ("h1", "h2"):
        cur = val
        sections[cur] = []
    else:
        sections[cur].append((kind, val))

d = docx.Document(TPL)
P = list(d.paragraphs)


def para_by_text(start):
    hits = [p for p in d.paragraphs if p.text.strip().startswith(start)]
    assert len(hits) == 1, f"{len(hits)} paragraphs start with {start!r}"
    return hits[0]


def set_runs(p, parts):
    """parts: list of (text, bold). Keep the first run's formatting for every new run."""
    proto = copy.deepcopy(p.runs[0]._r) if p.runs else None
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for text, bold in parts:
        if proto is not None:
            el = copy.deepcopy(proto)
            p._p.append(el)
            run = docx.text.run.Run(el, p)
            run.text = text
        else:
            run = p.add_run(text)
        run.bold = bold
        run.italic = False


LABEL = re.compile(r"^(RQ\d – [^.]+\.|Validation of the combination \(O5\)\.|Problem\.|Originality and significance\.|"
                   r"Quality and gaps\.|Applications and implications\.|Included:|Excluded:|Theoretical implications\.|"
                   r"Methodological implications\.|Practical implications\.|Plan to apply the research in practice\.)\s")


def parts_of(text):
    m = LABEL.match(text)
    if m:
        return [(m.group(1) + " ", True), (text[len(m.group(0)):], False)]
    return [(text, False)]


def replace_block(guides, texts):
    """Replace the guidance paragraphs `guides` by one paragraph per entry of `texts` (cloned from guides[0])."""
    first = guides[0]
    anchor = first._p
    for t in texts:
        el = copy.deepcopy(first._p)
        anchor.addprevious(el)
        set_runs(Paragraph(el, first._parent), t if isinstance(t, list) else parts_of(t))
    for g in guides:
        g._p.getparent().remove(g._p)


def guides_between(start_text, end_text):
    ps = list(d.paragraphs)
    i = next(k for k, p in enumerate(ps) if p.text.strip().startswith(start_text))
    j = next(k for k, p in enumerate(ps) if k > i and p.text.strip().startswith(end_text))
    return [p for p in ps[i + 1:j] if p.text.strip()]


# ---------------------------------------------------------------- header and cover table
set_runs(para_by_text("RESEARCH PROPOSAL TEMPLATE"), [("RESEARCH PROPOSAL", True)])
cover = d.tables[0]
vals = [TITLE, "Information Technology", "",
        "Bùi Anh Duy (SE193264, anhduy290605@gmail.com); Khổng Trần Gia Huy (SE180313, huyktgse180313@fpt.edu.vn)",
        "Nguyễn Thị Cẩm Hương (huongntc2@fpt.edu.vn); Trương Long (longt5@fpt.edu.vn)"]
for row, v in zip(cover.rows, vals):
    cell = row.cells[1]
    p = cell.paragraphs[0]
    if p.runs:
        set_runs(p, [(v, False)])
    else:
        p.add_run(v)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)

# ---------------------------------------------------------------- abstract and keywords
replace_block([para_by_text("Research objectives, research questions, research problems")], [ns["ABSTRACT"]])
kw = para_by_text("Key words:")
set_runs(kw, [("Key words: ", True), (ns["KEYWORDS"], False)])

# ---------------------------------------------------------------- 1.1, 1.2
replace_block(guides_between("1.1. Literature review", "1.2. The necessity"),
              [cite(v) for k, v in sections["1.1 Literature Review"]])
replace_block(guides_between("1.2. The necessity", "1.3. Feasibility"),
              [cite(v) for k, v in sections["1.2 Necessity of the Research"]])

# ---------------------------------------------------------------- 1.3 as the template's "+ Label: ..." items
feas = []
for k, v in sections["1.3 Feasibility of Research"]:
    lab, rest = v.split(". ", 1)
    feas.append([("+ " + lab + ": ", True), (cite(rest), False)])
replace_block(guides_between("1.3. Feasibility", "Research objectives"), feas)

# ---------------------------------------------------------------- 2-4, 6
replace_block(guides_between("Research objectives", "Research scope"), [cite(v) for k, v in sections["2. Research Objectives"]])
replace_block(guides_between("Research scope", "Approach and Method"), [cite(v) for k, v in sections["3. Research Scope"]])
replace_block(guides_between("Approach and Method", "Research plan"), [cite(v) for k, v in sections["4. Approach and Method"]])
replace_block(guides_between("Expected results", "References"), [cite(v) for k, v in sections["6. Expected Results"]])

# ---------------------------------------------------------------- 5 research plan table
intro = guides_between("Research plan", "Expected results")
plan = [v for k, v in sections["5. Research Plan"] if k == "plan"][0]
after_note = [v for k, v in sections["5. Research Plan"] if k == "p"]
set_runs(intro[0], [("The research plan, with expected outputs, schedule and person in charge, is given below.", False)])
tb = d.tables[1]
tmpl_row = tb.rows[1]._tr
for tr in [r._tr for r in tb.rows[1:]]:
    tr.getparent().remove(tr)
for i, (date, task, output, who) in enumerate(plan, 1):
    tr = copy.deepcopy(tmpl_row)
    tb._tbl.append(tr)
    cells = tb.rows[-1].cells
    for c, v in zip(cells, [str(i), date, task, output, who]):
        p = c.paragraphs[0]
        if p.runs:
            set_runs(p, [(v, False)])
        else:
            p.add_run(v)
        for extra in c.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
# sentence after the table
tbl_el = tb._tbl
note_el = copy.deepcopy(intro[0]._p)
tbl_el.addnext(note_el)
set_runs(Paragraph(note_el, intro[0]._parent), [(after_note[0], False)])

# ---------------------------------------------------------------- references
refs_head = para_by_text("References")
body_proto = para_by_text("The research plan, with expected outputs")
anchor = refs_head._p
for i, k in enumerate(order, 1):
    el = copy.deepcopy(body_proto._p)
    anchor.addnext(el)
    anchor = el
    q = Paragraph(el, refs_head._parent)
    set_runs(q, [(f"[{i}] {REFS[k]}", False)])
    q.paragraph_format.left_indent = Pt(22)
    q.paragraph_format.first_line_indent = Pt(-22)

# ---------------------------------------------------------------- appendix
app = para_by_text("Appendices (if any)")
set_runs(app, [("Appendices", True)])
el = copy.deepcopy(body_proto._p)
app._p.addnext(el)
ap_title = Paragraph(el, app._parent)
set_runs(ap_title, [("Appendix A. Preliminary results (from the public repository)", True)])
ta = d.add_table(rows=1 + len(ns["APPX"]), cols=2)
ta.style = d.tables[1].style
ta.cell(0, 0).text, ta.cell(0, 1).text = "Analysis", "Preliminary finding"
for j in range(2):
    ta.cell(0, j).paragraphs[0].runs[0].bold = True
for i, (a, b) in enumerate(ns["APPX"], 1):
    ta.cell(i, 0).text, ta.cell(i, 1).text = a, b
for row in ta.rows:
    for c in row.cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
ap_title._p.addnext(ta._tbl)
el = copy.deepcopy(body_proto._p)
ta._tbl.addnext(el)
set_runs(Paragraph(el, app._parent), [("These results come from three test windows; the combination was specified after the "
                                        "main analysis and is validated on unseen periods in the research plan.", False)])

# ---------------------------------------------------------------- commitment: template wording kept; title, disclosure, names
p = para_by_text("We are authors of the research paper titled")
t = p.text
assert "“...”" in t
set_runs(p, [(t.replace("“...”", f"“{TITLE}”"), False)])
hereby = para_by_text("Hereby, we would like to confirm")
el = copy.deepcopy(hereby._p)
hereby._p.addprevious(el)
set_runs(Paragraph(el, hereby._parent), [("Disclosure of AI use (rule 4): ", True),
                                          (ns["AI_DISCLOSURE"].replace("Disclosure of AI use (rule 4). ", "", 1), False)])
sig = d.tables[-1] if d.tables[-1].rows[0].cells[0].text.strip() == "Author 1" else [t for t in d.tables if t.rows[0].cells[0].text.strip() == "Author 1"][0]
names = ["Bùi Anh Duy", "Khổng Trần Gia Huy", ""]
for j, n in enumerate(names):
    c = sig.rows[1].cells[j]
    p0 = c.paragraphs[0]
    base = p0.text
    if n:
        el = copy.deepcopy(p0._p)
        c._tc.append(el)
        set_runs(Paragraph(el, c), [("\n\n" + n, False)])

# ---------------------------------------------------------------- remove the template's closing "Note" (author instructions)
ps = list(d.paragraphs)
i = next(k for k, q in enumerate(ps) if q.text.strip() == "Note:")
for q in ps[i:]:
    q._p.getparent().remove(q._p)

body = d.element.body
for el in reversed(list(body.iterchildren())):
    tag = el.tag.split("}")[1]
    if tag == "sectPr":
        continue
    if tag == "p" and not Paragraph(el, None).text.strip():
        body.remove(el)
        continue
    break

# 2-3. fonts in table cells and plan-table column widths
from docx.shared import Cm
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for q in c.paragraphs:
                for r in q.runs:
                    r.font.name = "Times New Roman"
                    r._r.get_or_add_rPr().get_or_add_rFonts().set(docx.oxml.ns.qn("w:eastAsia"), "Times New Roman")
                    if r.font.size is None:
                        r.font.size = Pt(11)
widths = [Cm(1.0), Cm(3.2), Cm(5.6), Cm(4.0), Cm(3.0)]
for row in d.tables[1].rows:
    for c, w in zip(row.cells, widths):
        c.width = w
for c, w in zip(d.tables[1].columns, widths):
    c.width = w

d.core_properties.title = TITLE
d.save(OUT)
print("written", OUT, "| references:", len(order))
