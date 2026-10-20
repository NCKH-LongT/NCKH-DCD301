"""Build the ICIIT 2027 submission (ACM single-column Word template) from 07_paper_draft/v3_draft.md and v3_tables.md.

The conference template is not stored in the repository: download template.docx from https://www.iciit.org/sub.html
and pass its path. All template styles are kept (Title_document, Authors, Abstract, Head1, Head2, PostHeadPara, Para,
TableCaption, TableCell, FigureCaption, Bib_entry, …); the "Authors' Information" form at the end of the template is
preserved for the authors to fill in.

Usage:
    python code/build_iciit_docx.py --template path/to/template.docx --out 08_final_submission/ICIIT2027_paper_v3.docx
"""
import argparse
import copy
import os
import re
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from f2d import config  # noqa: E402

DRAFT = os.path.join(config.ROOT, "07_paper_draft", "v3_draft.md")
TABLES = os.path.join(config.ROOT, "07_paper_draft", "v3_tables.md")
# table number in v3_tables.md -> number in the paper
TABLE_MAP = {2: 2, 3: 3, 5: 4, 6: 5, 7: 6}
LAB = {"empirical": "EMP", "tsb_nb": "TSB-NB", "ets": "ETS", "lgb_tweedie": "LGB-T", "lgb_conformal": "LGB-C",
       "hgb_quantile": "HGB-Q", "lgb_quantile": "LGB-Q"}
# column widths (cm) of the generated tables, keyed by their number in v3_tables.md; text width is 15.4 cm
WIDTHS = {2: [1.6, 2.3, 2.3, 2.3, 2.3, 2.3, 2.3],
          3: [1.3, 1.3, 1.9, 1.3, 1.1, 1.7, 1.7, 1.7, 1.7, 1.7],
          5: [1.3, 1.5, 3.2, 3.1, 3.1, 3.2],
          6: [4.6, 2.7, 2.7, 2.7, 2.7],
          7: [1.3, 1.5, 2.3, 2.5, 2.5, 2.5, 2.8]}
# tables whose three-window cells (main / second / third) are stacked on three lines
STACKED = {2, 7}
CCS = ("• Applied computing → Forecasting; Supply chain management • Computing methodologies → Machine learning; "
       "Cluster analysis")


def renumber(text):
    """Table 5/6/7 of the draft become Table 4/5/6 in the paper; Figure 2 becomes Figure 1."""
    for old, new in ((5, 4), (6, 5), (7, 6)):
        text = re.sub(rf"\bTable {old}\b", f"Table {new}", text)
    return re.sub(r"\bFigure 2\b", "Figure 1", text)


def add_runs(par, text, size=None):
    """Minimal inline markdown: **bold**, *italic*, `code`."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.italic = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1])
        else:
            r = par.add_run(tok)
        if size:
            r.font.size = Pt(size)


def md_table(lines):
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def parse_tables():
    """{number: (caption, rows, footnotes)} from v3_tables.md."""
    out, cur = {}, None
    for ln in open(TABLES, encoding="utf-8").read().split("\n"):
        m = re.match(r"\*\*Table (\d+)\.\*\*\s*(.*)", ln)
        if m:
            cur = int(m.group(1))
            out[cur] = [m.group(2), [], []]
        elif ln.startswith("**Table") or ln.startswith("(a)") or ln.startswith("# "):
            cur = None
        elif cur is not None:
            if ln.startswith("|"):
                out[cur][1].append(ln)
            elif ln.strip() and out[cur][1]:
                out[cur][2].append(ln.strip())
    return {k: (v[0], md_table(v[1]), v[2]) for k, v in out.items()}


def cell_text(text, compact):
    """Compact three-window cells (a / b / c -> a/b/c) and put confidence intervals on their own line."""
    if compact == "stack":
        text = text.replace(" / ", "\n")
    elif compact:
        text = text.replace(" / ", "/")
    return re.sub(r"\s+\[", "\n[", text)


def set_cell_margins(table, twips=40):
    """Small left/right cell margins so that narrow columns do not wrap."""
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def add_table(doc, caption, rows, notes=(), size=7.0, widths=None, compact=True):
    cap = doc.add_paragraph(style="TableCaption")
    add_runs(cap, caption)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_cell_margins(t)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            c = t.cell(i, j)
            p = c.paragraphs[0]
            p.style = doc.styles["TableCell"]
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            pf.first_line_indent = Pt(0)
            pf.left_indent = pf.right_indent = Pt(0)
            parts = (text.replace("**", "") if i == 0 else cell_text(text, compact)).split("\n")
            for k, part in enumerate(parts):
                if k:
                    p.add_run().add_break()
                add_runs(p, part, size)
            if i == 0:
                for r in p.runs:
                    r.bold = True
            if widths:
                c.width = Cm(widths[j])
    for n in notes:
        p = doc.add_paragraph(style="TableFootnote")
        add_runs(p, n, 7.0)
    return t


def make_figure(path):
    D = config.DEFAULT
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.0))
    for ax, d, tag in zip(axes, ("M5", "VN1"), ("(a) M5", "(b) VN1")):
        k = pd.read_csv(os.path.join(config.OUTPUTS, d, "kpi.csv"))
        k = k[(k.policy == "none") & (k.group == "all") & (k.L == D["L"]) & (k.H == D["H"]) & (k.q_liq == D["q_liq"])
              & (k.k == D["k_fixed"]) & k.model.isin(LAB)]
        for m in LAB:
            g = k[k.model == m].sort_values("tau")
            ax.plot(g.inventory_weeks, g.fill_rate, marker="o", ms=3, lw=1.2, label=LAB[m],
                    ls="-" if m.endswith("quantile") else "--")
        ax.set_xlabel("Average inventory (weeks of demand)", fontsize=8)
        ax.set_ylabel("Fill rate", fontsize=8)
        ax.set_title(tag, fontsize=9)
        ax.set_ylim(0.80, 1.0)
        ax.set_xlim(0.5, 6.0 if d == "VN1" else 3.5)
        ax.tick_params(labelsize=7)
        ax.grid(alpha=.3)
    axes[1].legend(fontsize=6.5, loc="lower right", ncol=2)
    fig.tight_layout()
    fig.savefig(path, dpi=300)
    plt.close(fig)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    doc = Document(a.template)
    body = doc.element.body
    # keep the "Authors' Information" form (and the final section properties); drop the rest of the sample content
    children = list(body.iterchildren())
    tail_from = None
    for i, ch in enumerate(children):
        if ch.tag == qn("w:p") and "Authors’ Information" in "".join(t.text or "" for t in ch.iter(qn("w:t"))):
            tail_from = i
    sect = children[-1]
    tail = [copy.deepcopy(c) for c in children[tail_from:-1]] if tail_from is not None else []
    for ch in children[:-1]:
        body.remove(ch)

    tabs = parse_tables()
    text = open(DRAFT, encoding="utf-8").read()
    lines = text.split("\n")
    title = lines[0].lstrip("# ").strip()

    doc.add_paragraph(title, style="Title_document")
    doc.add_paragraph("[AUTHOR NAMES TO BE ADDED]", style="Authors")
    doc.add_paragraph("[Affiliation, city, country, e-mail — to be added]", style="Affiliation")

    section, after_head, i = None, False, 1
    fig_path = os.path.join(os.path.dirname(os.path.abspath(a.out)), "fig1_tradeoff.png")

    def para(txt):
        nonlocal after_head
        p = doc.add_paragraph(style="PostHeadPara" if after_head else "Para")
        add_runs(p, renumber(txt))
        after_head = False
        return p

    while i < len(lines):
        ln = lines[i].rstrip()
        i += 1
        if not ln or ln.startswith(">"):
            continue
        if ln.startswith("## "):
            head = ln[3:].strip()
            section = head
            if head == "Abstract":
                continue
            if head == "References":
                doc.add_paragraph("REFERENCES", style="ReferenceHead")
                continue
            doc.add_paragraph(re.sub(r"^\d+\s+", "", head), style="Head1")
            after_head = True
            continue
        if ln.startswith("### "):
            doc.add_paragraph(re.sub(r"^\d+\.\d+\s+", "", ln[4:].strip()), style="Head2")
            after_head = True
            continue
        if section == "Abstract":
            if ln.startswith("**Keywords:**"):
                p = doc.add_paragraph(style="CCSDescription")
                p.add_run("CCS CONCEPTS ").bold = True
                p.add_run(CCS)
                p = doc.add_paragraph(style="KeyWords")
                p.add_run("Additional Keywords and Phrases: ").bold = True
                p.add_run(ln.replace("**Keywords:**", "").strip())
                doc.add_paragraph("ACM Reference Format:", style="ACMRefHead")
                doc.add_paragraph("[Author names to be added]. 2027. " + title + ". In Proceedings of the 12th International "
                                  "Conference on Intelligent Information Technology (ICIIT 2027). ACM, New York, NY, USA, 10 pages.",
                                  style="ACMRef")
            else:
                add_runs(doc.add_paragraph(style="Abstract"), ln)
            continue
        if section == "References":
            m = re.match(r"\[(\d+)\]\s+(.*)", ln)
            if m:
                add_runs(doc.add_paragraph(style="Bib_entry"), m.group(2))
            continue
        m = re.match(r"\*\[(Table|Figure) (\d+) here.*\]\*", ln)
        if m:
            kind, num = m.group(1), int(m.group(2))
            if kind == "Figure":
                make_figure(fig_path)
                p = doc.add_paragraph(style="Image")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(fig_path, width=Cm(15.0))
                doc.add_paragraph("Figure 1: Inventory (weeks of demand) against fill rate as the target service level τ varies "
                                  "from 0.5 to 0.99, main test window: (a) M5, (b) VN1. Curves further up and to the left are "
                                  "better. Solid lines: direct quantile models.", style="FigureCaption")
            else:
                cap, rows, notes = tabs[num]
                add_table(doc, f"Table {TABLE_MAP[num]}: " + cap, rows, notes, size=6.5, widths=WIDTHS[num],
                          compact="stack" if num in STACKED else True)
            after_head = False
            continue
        if ln.startswith("**Table 1.**"):
            cap = ln.replace("**Table 1.**", "").strip()
            tl = []
            while i < len(lines) and (lines[i].startswith("|") or not lines[i].strip()):
                if lines[i].startswith("|"):
                    tl.append(lines[i])
                elif tl:
                    break
                i += 1
            add_table(doc, "Table 1: " + cap, md_table(tl), size=7.0, widths=[4.4, 5.5, 5.5], compact=False)
            continue
        m = re.match(r"(\d+)\.\s+(.*)", ln)
        if m:
            para(f"({m.group(1)}) {m.group(2)}")
            continue
        if ln.startswith("- "):
            para(ln[2:])
            continue
        para(ln)

    if tail:
        doc.add_page_break()
    for el in tail:
        body.append(el)
    body.append(sect)
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    doc.save(a.out)
    print("written", a.out)


if __name__ == "__main__":
    main()
