"""Final file with the author block in the format the authors supplied (rev22 -> ICIIT2027_Paper.docx; 10 pages).

Format (from the sample the authors gave): names given-name first without diacritics (full syllable reversal, as in
"Binh Thuan Truong" for Trương Thuận Bình), superscript affiliation numbers, "and" before the last author, * on the
corresponding author; one line per affiliation starting with its superscript number; a last line with the
corresponding author's e-mail in a monospaced face. The ACM reference line uses the same names. Text otherwise
identical to rev22.

Usage:
    python code/revise_iciit_paper_final.py --src 08_final_submission/ICIIT2027_Paper_rev22.docx --out 08_final_submission/ICIIT2027_Paper.docx
"""
import argparse
import copy

import docx

AUTHORS = [("Duy Anh Bui", "1"), ("Huy Gia Tran Khong", "1"), ("Huong Cam Thi Nguyen", "1"), ("Long Truong", "1*")]
AFFILIATIONS = [("1", "Faculty of Software Engineering, FPT University, Ho Chi Minh City, Vietnam")]
CORRESPONDING_EMAIL = "longt5@fpt.edu.vn"
OLD_REF_NAMES = "Bùi Anh Duy, Khổng Trần Gia Huy, Nguyễn Thị Cẩm Hương, and Trương Long."
NEW_REF_NAMES = "Duy Anh Bui, Huy Gia Tran Khong, Huong Cam Thi Nguyen, and Long Truong."


def rebuild(p, parts, font=None):
    """parts: list of (text, superscript). New runs copy the formatting of the paragraph's first run."""
    proto = copy.deepcopy(p.runs[0]._r)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for text, sup in parts:
        el = copy.deepcopy(proto)
        p._p.append(el)
        run = docx.text.run.Run(el, p)
        run.text = text
        run.font.superscript = True if sup else None
        if font:
            run.font.name = font
            run._r.get_or_add_rPr().get_or_add_rFonts().set(docx.oxml.ns.qn("w:hAnsi"), font)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    d = docx.Document(a.src)
    names, aff = d.paragraphs[1], d.paragraphs[2]
    assert names.text.startswith("Bùi Anh Duy") and aff.text.startswith("FPT University")

    parts = []
    for k, (n, mark) in enumerate(AUTHORS):
        if k == len(AUTHORS) - 1:
            parts.append(("and ", False))
        parts.append((n, False))
        parts.append((mark, True))
        if k < len(AUTHORS) - 1:
            parts.append((", ", False))
    rebuild(names, parts)

    # one paragraph per affiliation, then the corresponding author's e-mail
    num, text = AFFILIATIONS[0]
    rebuild(aff, [(num, True), (" " + text, False)])
    last = aff
    for num, text in AFFILIATIONS[1:]:
        el = copy.deepcopy(aff._p)
        last._p.addnext(el)
        last = docx.text.paragraph.Paragraph(el, aff._parent)
        rebuild(last, [(num, True), (" " + text, False)])
    el = copy.deepcopy(aff._p)
    last._p.addnext(el)
    mail = docx.text.paragraph.Paragraph(el, aff._parent)
    rebuild(mail, [(CORRESPONDING_EMAIL, False)], font="Courier New")

    ref = [p for p in d.paragraphs if p.text.startswith("ACM Reference Format:")]
    assert len(ref) == 1 and OLD_REF_NAMES in ref[0].text
    for r in ref[0].runs:
        if OLD_REF_NAMES in r.text:
            r.text = r.text.replace(OLD_REF_NAMES, NEW_REF_NAMES)
            break
    else:
        raise SystemExit("reference names are split across runs; not changed")
    d.save(a.out)
    print("written", a.out, "|", names.text, "|", aff.text, "|", mail.text)


if __name__ == "__main__":
    main()
