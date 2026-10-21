"""Fourth review round on 08_final_submission/ICIIT2027_Paper_rev7.docx -> rev8 (must stay at 10 pages).

1  Default hyperparameters of the boosting models stated in Section 3.2 (values from f2d/models.py).
2  Censored demand: how censoring could enter training (EM / censored likelihood, censored quantile regression).
3  C2 caveat in the abstract and conclusion: the main evidence concerns the single methods.
4  Preprocessing of the eight clustering features (segment_series.py, class Prep) in Section 5.4.
5  Typography: Unicode minus sign in all tables (Table 3 used ASCII hyphens).

Usage:
    python code/revise_iciit_paper_r8.py --src 08_final_submission/ICIIT2027_Paper_rev7.docx --out 08_final_submission/ICIIT2027_Paper_rev8.docx
"""
import argparse
import re
import sys
import os

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from revise_iciit_paper_r7 import clone_after, find, replace_in, set_text  # noqa: E402

HPARAM = ("No hyperparameter search is performed. LGB-Q, LGB-T and LGB-C share one LightGBM configuration: learning "
          "rate 0.05, 63 leaves with no depth limit, at least 200 samples per leaf, L2 penalty 1, feature and row "
          "subsampling fractions of 0.8, and at most 1,000 boosting rounds with early stopping after 50 rounds. Each "
          "refit trains on up to 104 weekly origins and validates on the following 13; LGB-T and LGB-C use Tweedie loss "
          "with variance power 1.1. HGB-Q uses learning rate 0.1, at most 300 iterations, 63 leaves and at least 200 "
          "samples per leaf. Sensitivity checks appear in Section 5.8.")

CENSOR_TRAIN = (" Censoring could also be handled in training rather than only in evaluation: flagged weeks could be treated "
                "as missing or right-censored observations and estimated with an expectation–maximisation step or a "
                "censored likelihood, as in lost-sales demand estimation [28], with censored quantile regression [29] for "
                "the quantile learners, or with a survival model of the time between sales. Because the flags are "
                "uncertain, such models would need on-hand records or a sensitivity analysis over the flagging threshold; "
                "we leave them to future work.")

REF28 = "[28]  S. Nahmias. 1994. Demand estimation in lost sales inventory systems. Naval Research Logistics 41, 6, 739–757."
REF29 = "[29]  J. L. Powell. 1986. Censored regression quantiles. Journal of Econometrics 32, 1, 143–155."

CLUSTER_PREP = (" ADI, CV² and time since sale are log-transformed; all eight features are then centred on their medians, "
                "divided by their interquartile ranges and clipped at ±4, with these statistics taken from the fitting "
                "window.")


def unicode_minus(doc):
    """ASCII hyphen directly before a digit (and not part of a word or range) -> U+2212 in every table cell."""
    n = 0
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        new = re.sub(r"(?<![A-Za-z0-9.])-(?=\d)", "−", r.text)
                        if new != r.text:
                            n += r.text.count("-") - new.count("-")
                            r.text = new
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    d = docx.Document(a.src)

    # 1 – hyperparameters in Section 3.2
    replace_in(d, "An additional combination, C2, specified post hoc",
               "No hyperparameter search is performed; sensitivity checks appear in Section 5.8.", HPARAM)

    # 2 – censoring in training
    p = find(d, "Censored demand. Both datasets record sales")
    set_text(p, p.text + CENSOR_TRAIN)
    q = clone_after(find(d, "[27]  O. M. Nabil"), REF28)
    clone_after(q, REF29)

    # 3 – C2 caveat in abstract and conclusion
    replace_in(d, "The forecast with the lowest error is not necessarily",
               "This combination was selected after inspection of the results and has not been independently validated.",
               "This combination was selected after inspection of the results and has not been independently validated, "
               "so our main evidence concerns the single methods.")
    replace_in(d, "The VN1 results do not identify a single leading method",
               "Because that pair was chosen after inspection of the results, the combination needs an independent evaluation.",
               "Because that pair was chosen after inspection of the results, this finding is exploratory and needs an "
               "independent evaluation; the conclusions about single methods above do not depend on it.")

    # 4 – clustering preprocessing in Section 5.4
    replace_in(d, "We next ask whether information available before a window",
               "recent relative level and peak size.", "recent relative level and peak size." + CLUSTER_PREP)

    # 5 – minus signs in tables
    print("table minus signs replaced:", unicode_minus(d))
    d.save(a.out)
    print("written", a.out)


if __name__ == "__main__":
    main()
