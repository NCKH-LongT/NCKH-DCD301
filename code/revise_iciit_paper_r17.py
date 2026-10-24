"""rev16 -> rev17 (must stay at 10 pages).

P0  Table 7 and its reading: target vs achieved fill; which comparisons are like-for-like (achieved fills within
    0.5 points of LGB-Q's) and which are not; no superiority claim where service differs. Flags are computed from
    code/outputs/robustness/rolling_calibration.csv.
P1  Section 4 matches the design: three windows for M5 and VN1, one for the footwear panel.
P1  Discussion states why moving from ex-post to prospective evaluation matters and the limits of transferring a ranking.
P2  "gross margin" -> "markup" (the formula uses price = (1 + m) × cost); reproducibility sentence for the footwear
    panel when --repro-ok is given (forecasts rerun from the public files match the cached ones).

Usage:
    python code/revise_iciit_paper_r17.py --src 08_final_submission/ICIIT2027_Paper_rev16.docx --out 08_final_submission/ICIIT2027_Paper_rev17.docx [--repro-ok]
"""
import argparse
import os
import sys

import docx
import pandas as pd
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from revise_iciit_paper_r7 import find, replace_in, set_text  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
COLS = [("M5", "third → second"), ("M5", "second → main"), ("VN1", "third → second"), ("VN1", "second → main"),
        ("VNF", "main (τ0 = target)")]
ROWS = ["LGB-Q", "HGB-Q", "TSB-NB", "C2", "E7-med"]

CAPTION = ("Table 7. Past-only service calibration (τ updated weekly, η = 0.5); target fill 0.92 throughout, Fill = "
           "achieved fill. ΔI = inventory relative to LGB-Q (%), 95% cluster-bootstrap intervals (M5 items, VN1 warehouses, "
           "footwear styles); LGB-Q rows: weeks of demand. † = fills differ by ≥ 0.5 points, not like-for-like.")

READING = ("Table 7 separates the target (0.92) from the achieved fill (0.906–0.926 on M5, 0.906–0.940 on VN1, "
           "0.865–0.893 on the footwear panel). Only 13 of 20 comparisons have achieved fills within 0.5 points of LGB-Q's, "
           "none on the footwear panel; two of these show a significant saving (E7-med on M5, 0.9%; C2 on VN1, 7.4%), eight "
           "need significantly more stock and three are inconclusive, so no method is shown to beat LGB-Q where service "
           "differs.")

DISC_ADD = (" Matching realised fill sets τ with the test outcomes, whereas a deployed system must fix it beforehand; "
            "on VN1 and the footwear panel the achieved fill then misses the target by up to 5.5 points, so retrospective "
            "savings do not become reliable stock reductions, and a ranking should not be transferred without a "
            "prospective check.")

VNF_SHORT = ("Independent panel. To check the combinations on data not used to specify them, we add the "
             "retail-channel sales of a Vietnamese footwear chain, public on Kaggle "
             "(tienanh2003/sales-and-inventory-snapshot-data): 1,001 style–colour series over 82 weeks (January 2022 to July "
             "2023), excluding returns and three partial weeks. With the eligibility rule above, 909 series enter one 26-week "
             "test window from 30 January 2023. Forecasts use the same pipeline, 21 features, refit schedule and policy; "
             "rerunning it on the processed sales file reproduces them exactly. Only one window fits the short history, "
             "month-end stock snapshots are unused, and single-method results on this panel existed before, but were not "
             "used for, specifying the combinations.")

REPRO = (" Rerunning the forecasting pipeline on the processed sales file reproduces the cached forecasts "
         "exactly.")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--repro-ok", action="store_true")
    a = ap.parse_args()
    d = docx.Document(a.src)

    # P0 – Table 7: caption, † flags, reading in Section 5.8
    set_text(find(d, "Table 7. Past-only service calibration"), CAPTION)
    r = pd.read_csv(os.path.join(ROOT, "code", "outputs", "robustness", "rolling_calibration.csv"))
    r = r[(r.eta == 0.5) & (r.target == 0.92)]
    t = d.tables[6]
    assert t.cell(0, 0).text.strip() == "Method"
    n_flag = 0
    for i, m in enumerate(ROWS):
        if m == "LGB-Q":
            continue
        for j, (ds, tr) in enumerate(COLS):
            row = r[(r.dataset == ds) & (r.transition == tr) & (r.method == m)].iloc[0]
            if abs(row.d_fill) >= 0.005:
                p = t.cell(2 + i, 2 + 2 * j).paragraphs[0]
                first_t = p.runs[0]._r.findall(qn("w:t"))[0]      # the ΔI value, before the line break
                first_t.text = first_t.text + "†"
                n_flag += 1
    assert n_flag == 7, n_flag
    replace_in(d, "Prospective τ, resampling and a third dataset.",
               "Table 7 reports achieved fill and inventory together under this past-only rule; we read an inventory "
               "difference as a saving only where fills differ by less than 0.5 points. In that sense E7-med never needs "
               "significantly less stock than LGB-Q on VN1 or on the footwear panel, and C2 does so once.", READING)

    # P1 – Section 4
    replace_in(d, "For each dataset, the main window comprises the final 26 weeks;",
               "For each dataset, the main window comprises the final 26 weeks; two earlier windows are formed by removing "
               "the last 26 and 52 weeks, respectively.",
               "For M5 and VN1, the main window comprises the final 26 weeks; two earlier windows are formed by removing "
               "the last 26 and 52 weeks, respectively. The footwear panel has only its final 26-week window.")
    replace_in(d, "For M5 and VN1, the main window", "Training is repeated for all three windows.",
               "Training is repeated for every window.")

    # P1 – significance of the ex-post -> prospective shift
    p = find(d, "The comparison metric changes what counts as a good forecast.")
    set_text(p, p.text + DISC_ADD)

    # P2 – markup, reproducibility
    replace_in(d, "Liquidation. For each policy we simulate", "gross margin m (selling price (1 + m) times cost)",
               "markup m (selling price (1 + m) times cost)")
    replace_in(d, "Liquidation. For each policy we simulate", "gross margins of 30–100%", "markups of 30–100%")
    replace_in(d, "Table 6. Changes relative to no liquidation", "holding-cost and margin assumptions",
               "holding-cost and markup assumptions")
    # compensating cuts (repeated elsewhere) and the window count in Limitations
    replace_in(d, "On VN1, the retrospective demand split",
               " In our experiment, the ranking also changes with the test period; no comparable reversal appears on M5 at "
               "the fill rates considered.", "")
    replace_in(d, "The VN1 rankings are less clear at series level.",
               " LGB-Q's aggregate wQL advantage thus does not extend uniformly to individual VN1 series.", "")
    replace_in(d, "Limitations. Recorded sales stand in for demand",
               " Lead times of one, two and four weeks were examined in all three windows, without HGB-Q in the lead-time "
               "checks.", "")
    replace_in(d, "Limitations. Recorded sales stand in for demand",
               "and each dataset supplies only three evaluation windows.",
               "and M5 and VN1 supply three evaluation windows, the footwear panel one.")
    replace_in(d, "Limitations. Recorded sales stand in for demand",
               "Chronos-2 [22] was run zero-shot only; N-HiTS is an additional neural baseline (Section 5.8), not one of the "
               "seven methods in the main benchmark.", "Chronos-2 [22] ran zero-shot only.")
    if a.repro_ok:
        set_text(find(d, "Independent panel."), VNF_SHORT)
    replace_in(d, "We tested whether specific differences between M5 and VN1",
               "Removing price variables from VN1, where prices are available in fewer than half of the weeks, or",
               "Removing VN1's sparse price variables or")
    d.save(a.out)
    print("written", a.out, "| flagged cells:", n_flag)


if __name__ == "__main__":
    main()
