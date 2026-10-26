"""Final improvement plan (phase5_workflow.md, section 10) in rev18 -> rev19 (must stay at 10 pages).

Item 1  Two research questions; contributions as two core findings plus supporting results; Related Work states how
        the design differs from the closest studies (07_paper_draft/contribution_matrix.md, checked against full texts).
Item 2  Table 7 intervals from the full-policy cluster bootstrap (τ path recomputed in each of 2,000 paired replicates;
        code/outputs/robustness/full_policy_bootstrap.csv); point values unchanged (identical to rev18).
Item 4  Section 5.8 counts and the fill-equivalence sensitivity from code/outputs/robustness/service_tradeoff.md.
Item 3  Warehouse-level ranking-reversal result (code/outputs/robustness/ranking_reversal.md) in the Discussion.
Compensating cuts keep the paper at 10 pages.

Usage:
    python code/revise_iciit_paper_r19.py --src 08_final_submission/ICIIT2027_Paper_rev18.docx --out 08_final_submission/ICIIT2027_Paper_rev19.docx
"""
import argparse
import os
import sys

import docx
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from revise_iciit_paper_r7 import find, replace_in, set_text  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
COLS = [("M5", "third → second"), ("M5", "second → main"), ("VN1", "third → second"), ("VN1", "second → main"),
        ("VNF", "main (τ0 = target)")]
ROWS = ["LGB-Q", "HGB-Q", "TSB-NB", "C2", "E7-med"]

LEAD = ("We hold the inventory policy fixed and vary only its demand forecasts, evaluating seven methods on every "
        "eligible series of M5 and VN1 over three 26-week windows each. We ask two questions. RQ1: do rankings by forecast "
        "accuracy and by the inventory needed for a given realised fill rate agree across retail settings and periods? "
        "RQ2: do inventory savings measured at matched realised fill survive deployment-style evaluation, with service "
        "calibrated from past weeks only and an independent panel? Our contributions are as follows.")

BULLETS = [
    "•  Ranking agreement depends on the setting and the period (RQ1). On M5, LGB-Q is both the most accurate method and "
    "needs 3–11% less stock than the point-forecast, conformal, TSB-NB and exponential-smoothing alternatives in every "
    "window; on VN1 the leading method changes between windows, and neither demand classes, clusters nor pre-test "
    "warehouse features select it reliably in advance.",
    "•  Retrospective savings are not prospective evidence (RQ2). The 7–8% VN1 saving of a pre-specified seven-method "
    "median at matched realised fill disappears under past-only calibration, full-policy cluster bootstraps and an "
    "independent footwear panel, whereas the M5 ranking holds; achieved fill is reported beside every inventory "
    "comparison.",
    "•  Supporting results. Checks on lead time, hyperparameters, censored demand and computing cost show which findings "
    "are stable, a break-even salvage ratio states when liquidation pays, and code and results are public.",
]

RELATED = ("We retain eligible series across demand classes, use two retail settings with three test windows plus an "
           "independent third panel, and include a liquidation decision. We also estimate quantiles directly rather than "
           "deriving all of them from residuals. Inventory is compared at the same realised fill rate, because a common "
           "nominal quantile level does not give all methods the same achieved service [7]. Neither multi-period stock "
           "simulation nor decision-based forecast evaluation is new [5, 7, 9, 14]. Unlike [7], which evaluates one "
           "365-day period at product–state level, and [5], which tunes combinations for a single-period newsvendor on "
           "one 28-day holdout, we compare rankings across settings and periods and test whether retrospective savings "
           "survive prospective evaluation.")

READING = ("Table 7 separates the target (0.92) from the achieved fill (0.906–0.926 on M5, 0.906–0.940 on VN1, "
           "0.865–0.893 on the footwear panel); LGB-Q meets the target within 0.5 points in one of five cases. Its "
           "intervals rerun the past-only policy, τ path included, in 2,000 paired cluster replicates and are wider "
           "than with τ held fixed, most on VN1. Of the 13 comparisons whose achieved fills are within "
           "0.5 points of LGB-Q's (none on the footwear panel), one shows a significant saving (C2 on VN1, 7.4%), five need "
           "significantly more stock and seven are inconclusive; with thresholds of 0.25 or 1.0 points, 10 or 17 qualify "
           "and one or two show a saving. No method is therefore shown to beat LGB-Q where service differs.")

CAPTION = ("Table 7. Past-only service calibration (τ updated weekly, η = 0.5); target fill 0.92 throughout, Fill = "
           "achieved fill. ΔI = inventory relative to LGB-Q (%), 95% full-policy cluster-bootstrap intervals (τ path "
           "recomputed per replicate; M5 items, VN1 warehouses, footwear styles); LGB-Q rows: weeks of demand. "
           "† = fills differ by ≥ 0.5 points, not like-for-like.")

REVERSAL = ("At warehouse level, pre-test features carry some information on the reversals: fitted on an earlier "
            "window, they rank the TSB-NB/LGB-Q inventory ratio of the next window with Spearman ρ of 0.21–0.27 "
            "(permutation p ≤ 0.046), but assigning TSB-NB to the warehouses they favour changes VN1 inventory by −9.4%, "
            "−0.2% and +2.9% across three train–test splits, so the signal does not yield a reliable selection rule.")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    d = docx.Document(a.src)

    # item 1 – RQs, contributions, related work
    set_text(find(d, "We hold the inventory policy fixed"), LEAD)
    bullets = [p for p in d.paragraphs if p.text.startswith("•  ")]
    assert len(bullets) == 4
    for p, text in zip(bullets[:3], BULLETS):
        set_text(p, text)
    bullets[3]._p.getparent().remove(bullets[3]._p)
    set_text(find(d, "We retain eligible series across demand classes"), RELATED)

    # item 2 – Table 7 intervals from the full-policy bootstrap
    fp = pd.read_csv(os.path.join(ROOT, "code", "outputs", "robustness", "full_policy_bootstrap.csv"))
    sg = lambda v: f"{100 * v:+.1f}".replace("-", "−")
    t = d.tables[6]
    n = 0
    for i, m in enumerate(ROWS):
        if m == "LGB-Q":
            continue
        for j, (ds, tr) in enumerate(COLS):
            x = fp[(fp.dataset == ds) & (fp.transition == tr) & (fp.method == m)].iloc[0]
            runs = t.cell(2 + i, 2 + 2 * j).paragraphs[0].runs
            assert runs[0].text.replace("†", "").strip().startswith(sg(x.d_inv)), (m, ds, tr, runs[0].text, sg(x.d_inv))
            runs[1].text = f"[{sg(x.dinv_lo)},{sg(x.dinv_hi)}]"
            n += 1
    assert n == 20
    set_text(find(d, "Table 7. Past-only service calibration"), CAPTION)

    # item 4 – reading of Table 7 in Section 5.8
    p = find(d, "Prospective τ, resampling and a third dataset.")
    txt = p.text
    i0 = txt.index("Table 7 separates the target")
    i1 = txt.index("Resampling items, stores or blocks of weeks")
    set_text(p, txt[:i0] + READING + " " + txt[i1:])

    # item 3 – ranking reversal in the Discussion (replaces the reference to the repository classifier)
    replace_in(d, "On VN1, the retrospective demand split",
               "A classifier based on earlier features recovers only part of that hypothetical saving and remains unstable "
               "(repository).", REVERSAL)

    # compensating cuts
    replace_in(d, "Among the combinations, E7-med has the smallest maximum regret",
               " When a lost sale costs at most 200 weeks of holding, C2's cost is within 3.1% of the cheapest single model "
               "in each run. At higher penalties LGB-Q is cheapest.", "")
    replace_in(d, "Censored demand. Both datasets record sales",
               ", about twice as often in intermittent and lumpy as in smooth series;", ";")
    set_text(find(d, "Toward deployment."),
             "Toward deployment. Because neither combination passed the prospective checks of Section 5.8, LGB-Q is the "
             "safer default, and retrospective savings alone do not justify a switch. A candidate should be frozen and run "
             "in shadow mode beside it for several review cycles, with achieved fill and upper-quantile coverage monitored, "
             "where TSB-NB under-covers.")
    replace_in(d, "Neural baseline. We evaluated", " Including N-HiTS leaves the post hoc C2's maximum regret at 0.9%.", "")
    replace_in(d, "How the quantiles are constructed also matters",
               "Residual-based quantiles in [7] and empirical quantiles in [23] may explain part of the disagreement with "
               "those studies, which do not isolate this effect.",
               "Different quantile construction may explain part of the disagreement with [7] and [23].")
    replace_in(d, "We tested whether specific differences between M5 and VN1",
               "The tested covariate and preprocessing differences do not account for the results; other dataset differences "
               "remain untested.", "Other dataset differences remain untested.")
    replace_in(d, "The VN1 results do not identify a single leading method",
               " TSB-NB uses less stock on retrospectively identified stable-demand series, while LGB-Q achieves a 0.92 fill "
               "rate on growing series in two windows where TSB-NB cannot.", "")
    d.save(a.out)
    print("written", a.out, "| Table 7 intervals replaced:", n)


if __name__ == "__main__":
    main()
